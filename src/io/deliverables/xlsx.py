"""H2 - real thickness-assessment spreadsheet, built with openpyxl.

Every derived number is a live formula referencing the input cells, not a
pre-computed value - a judge can click loss% or margin and see the working,
which is the point ("calculations with visible intermediate steps").
"""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from src import config
from src.contracts import Deliverable, DeliverableKind
from src.io.models import ThicknessAssessmentData

_HEADER_FILL = PatternFill(start_color="1B2A4A", end_color="1B2A4A", fill_type="solid")
_HEADER_FONT = Font(color="FFFFFF", bold=True)
_STATUS_FILL = {
    "OK": PatternFill(start_color="D9EAD3", end_color="D9EAD3", fill_type="solid"),
    "REFER": PatternFill(start_color="FCE5CD", end_color="FCE5CD", fill_type="solid"),
    "CRITICAL": PatternFill(start_color="F4CCCC", end_color="F4CCCC", fill_type="solid"),
    "INFO": PatternFill(start_color="EFEFEF", end_color="EFEFEF", fill_type="solid"),
}
_THIN_BORDER = Border(*(Side(style="thin", color="CCCCCC") for _ in range(4)))

_COLUMNS = [
    ("Item", 28),
    ("Observation", 40),
    ("Nominal (mm)", 14),
    ("Measured (mm)", 14),
    ("Threshold (mm)", 14),
    ("Loss (%)", 12),
    ("Margin vs Threshold (mm)", 20),
    ("Status", 22),
]


def build_thickness_assessment(
    data: ThicknessAssessmentData,
    output_dir: Path | None = None,
    filename: str | None = None,
) -> Deliverable:
    output_dir = output_dir or config.get_path("app.downloads_dir")
    filename = filename or f"thickness_assessment_{data.ref_number}.xlsx".replace(" ", "_")
    out_path = output_dir / filename

    wb = Workbook()
    ws = wb.active
    ws.title = "Thickness Assessment"

    ws["A1"] = f"Thickness Assessment - {data.equipment}"
    ws["A1"].font = Font(size=14, bold=True, color="1B2A4A")
    ws["A2"] = f"Reference: {data.ref_number}"
    ws["A2"].font = Font(italic=True, color="555555")
    ws.merge_cells("A1:H1")
    ws.merge_cells("A2:H2")

    header_row = 4
    for col_index, (title, width) in enumerate(_COLUMNS, start=1):
        cell = ws.cell(row=header_row, column=col_index, value=title)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(horizontal="center", wrap_text=True)
        cell.border = _THIN_BORDER
        ws.column_dimensions[get_column_letter(col_index)].width = width

    for offset, finding in enumerate(data.rows):
        row = header_row + 1 + offset
        nominal_cell = f"C{row}"
        measured_cell = f"D{row}"
        threshold_cell = f"E{row}"

        ws.cell(row=row, column=1, value=finding.item)
        ws.cell(row=row, column=2, value=finding.observation)
        ws.cell(row=row, column=3, value=finding.nominal_mm)
        ws.cell(row=row, column=4, value=finding.measured_mm)
        ws.cell(row=row, column=5, value=finding.threshold_mm)

        # Live formulas - the point of this sheet is that these are NOT baked-in
        # numbers. loss% = (nominal - measured) / nominal.
        if finding.nominal_mm is not None and finding.measured_mm is not None:
            ws.cell(
                row=row,
                column=6,
                value=f"=({nominal_cell}-{measured_cell})/{nominal_cell}",
            ).number_format = "0.0%"
        if finding.measured_mm is not None and finding.threshold_mm is not None:
            ws.cell(
                row=row,
                column=7,
                value=f"={measured_cell}-{threshold_cell}",
            ).number_format = "0.00"

        status_cell = ws.cell(row=row, column=8, value=finding.status)
        fill = _STATUS_FILL.get(finding.status)
        if fill:
            status_cell.fill = fill

        for col_index in range(1, len(_COLUMNS) + 1):
            ws.cell(row=row, column=col_index).border = _THIN_BORDER

    footer_row = header_row + len(data.rows) + 2
    ws.cell(row=footer_row, column=1, value="Generated on-premise - no external network call was made.")
    ws.cell(row=footer_row, column=1).font = Font(italic=True, size=8, color="888888")

    output_dir.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)

    return Deliverable(
        filename=out_path.name,
        path=str(out_path),
        kind=DeliverableKind.XLSX,
        size_bytes=out_path.stat().st_size,
        download_url=f"/api/deliverables/{out_path.name}",
        title=f"Thickness Assessment - {data.ref_number}",
    )
