"""H2 - real Word approval note, built with python-docx.

Looks like an actual industrial document: header block, reference number,
equipment/inspection metadata, a findings table, a recommendation paragraph,
source citations, and a signature block. Not a text dump.
"""

from __future__ import annotations

from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

from src import config
from src.contracts import Deliverable, DeliverableKind
from src.io.models import ApprovalNoteData

_NAVY = RGBColor(0x1B, 0x2A, 0x4A)
_GREY = RGBColor(0x55, 0x55, 0x55)

_STATUS_LABEL = {
    "OK": "Within limits",
    "REFER": "Refer for engineering review",
    "CRITICAL": "Immediate action required",
    "INFO": "Informational",
}


def _set_cell_shading(cell, hex_color: str) -> None:
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement

    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), hex_color)
    cell._tc.get_or_add_tcPr().append(shading)


def build_approval_note(
    data: ApprovalNoteData,
    output_dir: Path | None = None,
    filename: str | None = None,
) -> Deliverable:
    output_dir = output_dir or config.get_path("app.downloads_dir")
    filename = filename or f"approval_note_{data.ref_number}.docx".replace(" ", "_")
    out_path = output_dir / filename

    doc = Document()

    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    header = doc.add_paragraph()
    header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = header.add_run(data.prepared_for.upper())
    run.bold = True
    run.font.size = Pt(16)
    run.font.color.rgb = _NAVY

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub_run = subtitle.add_run("EQUIPMENT APPROVAL NOTE")
    sub_run.font.size = Pt(12)
    sub_run.font.color.rgb = _GREY
    sub_run.italic = True

    doc.add_paragraph()

    meta_table = doc.add_table(rows=4, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.LEFT
    meta_rows = [
        ("Reference No.", data.ref_number),
        ("Equipment", data.equipment),
        ("Inspection Date", data.inspection_date),
        ("Inspector", data.inspector),
    ]
    for row, (label, value) in zip(meta_table.rows, meta_rows):
        row.cells[0].text = label
        row.cells[0].paragraphs[0].runs[0].bold = True
        row.cells[1].text = value

    doc.add_paragraph()

    findings_heading = doc.add_paragraph()
    findings_run = findings_heading.add_run("FINDINGS")
    findings_run.bold = True
    findings_run.font.size = Pt(13)
    findings_run.font.color.rgb = _NAVY

    table = doc.add_table(rows=1, cols=5)
    table.style = "Light Grid Accent 1"
    header_cells = table.rows[0].cells
    for cell, text in zip(header_cells, ["Item", "Observation", "Measured", "Threshold", "Status"]):
        cell.text = text
        cell.paragraphs[0].runs[0].bold = True
        _set_cell_shading(cell, "1B2A4A")
        cell.paragraphs[0].runs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    status_fill = {"OK": "D9EAD3", "REFER": "FCE5CD", "CRITICAL": "F4CCCC", "INFO": "EFEFEF"}
    for finding in data.findings:
        row = table.add_row()
        cells = row.cells
        cells[0].text = finding.item
        cells[1].text = finding.observation
        cells[2].text = (
            f"{finding.measured_mm:.1f} mm" if finding.measured_mm is not None else "-"
        )
        cells[3].text = (
            f"{finding.threshold_mm:.1f} mm" if finding.threshold_mm is not None else "-"
        )
        cells[4].text = _STATUS_LABEL.get(finding.status, finding.status)
        _set_cell_shading(cells[4], status_fill.get(finding.status, "FFFFFF"))

    doc.add_paragraph()

    rec_heading = doc.add_paragraph()
    rec_run = rec_heading.add_run("RECOMMENDATION")
    rec_run.bold = True
    rec_run.font.size = Pt(13)
    rec_run.font.color.rgb = _NAVY
    doc.add_paragraph(data.recommendation)

    if data.sources:
        doc.add_paragraph()
        src_heading = doc.add_paragraph()
        src_run = src_heading.add_run("SOURCES CITED")
        src_run.bold = True
        src_run.font.size = Pt(13)
        src_run.font.color.rgb = _NAVY
        for source in data.sources:
            para = doc.add_paragraph(style="List Bullet")
            label = Path(source.source_path).name
            para.add_run(f"{label}, page {source.page}: ").bold = True
            para.add_run(source.snippet)

    doc.add_paragraph()
    doc.add_paragraph()
    sig_table = doc.add_table(rows=2, cols=2)
    sig_table.rows[0].cells[0].text = "Inspector: " + data.inspector
    sig_table.rows[0].cells[1].text = "Reviewing Engineer: " + (data.reviewing_engineer or "____________________")
    sig_table.rows[1].cells[0].text = "Signature: ____________________"
    sig_table.rows[1].cells[1].text = "Date: " + date.today().isoformat()

    footer = doc.add_paragraph()
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_run = footer.add_run(
        "Generated on-premise by the Sovereign AI Workbench - no external network call was made."
    )
    footer_run.font.size = Pt(8)
    footer_run.font.color.rgb = _GREY
    footer_run.italic = True

    output_dir.mkdir(parents=True, exist_ok=True)
    doc.save(out_path)

    return Deliverable(
        filename=out_path.name,
        path=str(out_path),
        kind=DeliverableKind.DOCX,
        size_bytes=out_path.stat().st_size,
        download_url=f"/api/deliverables/{out_path.name}",
        title=f"Approval Note - {data.ref_number}",
    )
