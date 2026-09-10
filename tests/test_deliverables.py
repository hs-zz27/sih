"""H2 - deliverable generators must produce real, openable Office files."""

from __future__ import annotations

from docx import Document as DocxReader
from openpyxl import load_workbook

from src.contracts import SourceCitation
from src.io.deliverables.docx import build_approval_note
from src.io.deliverables.xlsx import build_thickness_assessment
from src.io.models import ApprovalNoteData, FindingRow, ThicknessAssessmentData

FINDINGS = [
    FindingRow(
        item="Shell course 2, grid C4",
        observation="Wall thickness below the SOP-114 retirement threshold.",
        measured_mm=7.1,
        nominal_mm=9.5,
        threshold_mm=7.6,
        status="REFER",
    ),
    FindingRow(
        item="Tube bundle, north header",
        observation="3 tubes plugged, no new leakage observed.",
        status="OK",
    ),
]


def test_approval_note_is_a_real_docx(tmp_path):
    data = ApprovalNoteData(
        ref_number="INSP-2026-0412",
        equipment="Shell-and-tube heat exchanger E-4102",
        inspection_date="2026-04-12",
        inspector="R. Nayak",
        findings=FINDINGS,
        recommendation="Refer to engineering for a fitness-for-service assessment.",
        sources=[
            SourceCitation(
                document_id="doc_1",
                source_path="data/corpus/SOP-114_heat_exchanger_inspection_scan.jpg",
                page=1,
                snippet="Referral threshold: 80% of nominal wall thickness.",
            )
        ],
    )

    deliverable = build_approval_note(data, output_dir=tmp_path)

    assert deliverable.path.endswith(".docx")
    reader = DocxReader(deliverable.path)
    full_text = "\n".join(p.text for p in reader.paragraphs)
    full_text += "\n" + "\n".join(
        cell.text for table in reader.tables for row in table.rows for cell in row.cells
    )
    assert "INSP-2026-0412" in full_text
    assert "R. Nayak" in full_text
    assert any("REFER" in cell.text or "Refer for engineering review" in cell.text
                for table in reader.tables for row in table.rows for cell in row.cells)


def test_thickness_assessment_has_live_formulas(tmp_path):
    data = ThicknessAssessmentData(
        ref_number="INSP-2026-0412",
        equipment="E-4102",
        rows=FINDINGS,
    )

    deliverable = build_thickness_assessment(data, output_dir=tmp_path)

    assert deliverable.path.endswith(".xlsx")
    wb = load_workbook(deliverable.path)
    ws = wb.active

    loss_formula = ws["F5"].value
    margin_formula = ws["G5"].value
    assert isinstance(loss_formula, str) and loss_formula.startswith("=")
    assert isinstance(margin_formula, str) and margin_formula.startswith("=")
    assert "C5" in loss_formula and "D5" in loss_formula
